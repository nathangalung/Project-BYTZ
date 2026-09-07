"""Declared type versus actual bytes.

Nothing on the upload path ever looked at the content. The presign now allows
only a fixed set of content types, but the type is still what the caller says
it is, so a caller can store HTML under application/pdf. Serving is handled at
the edge (nosniff plus Content-Disposition); this is the parser's own check, so
a mislabelled file is refused rather than fed to the model as garbage.
"""

import io

import pytest

from app.services.cv_parser import detect_signature, signature_matches


def _pptx(text: str = "hello") -> bytes:
    from pptx import Presentation

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


class TestDetectSignature:
    def test_reads_a_pdf(self):
        assert detect_signature(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n") == "pdf"

    def test_reads_a_zip_container(self):
        # docx and pptx are both zip archives, indistinguishable at byte 0.
        assert detect_signature(_pptx()) == "zip"

    def test_reads_legacy_ole(self):
        assert detect_signature(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest") == "ole"

    def test_calls_anything_else_unknown(self):
        assert detect_signature(b"<html><body>hi</body></html>") == "unknown"

    def test_handles_an_empty_body(self):
        assert detect_signature(b"") == "unknown"

    def test_handles_a_body_shorter_than_any_signature(self):
        assert detect_signature(b"%P") == "unknown"


class TestSignatureMatches:
    @pytest.mark.parametrize(
        "declared,body",
        [
            ("pdf", b"%PDF-1.4 rest"),
            ("docx", _pptx()),
            ("pptx", _pptx()),
            ("doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1x"),
        ],
    )
    def test_accepts_bytes_that_match(self, declared, body):
        assert signature_matches(declared, body) is True

    def test_refuses_html_claiming_to_be_a_pdf(self):
        assert signature_matches("pdf", b"<html><script>alert(1)</script></html>") is False

    def test_refuses_a_pdf_claiming_to_be_a_docx(self):
        assert signature_matches("docx", b"%PDF-1.4 rest") is False

    def test_refuses_a_zip_claiming_to_be_a_pdf(self):
        assert signature_matches("pdf", _pptx()) is False

    def test_allows_a_type_it_has_no_signature_for(self):
        # txt and md have no magic bytes; refusing them would break uploads
        # that are legitimately plain text.
        assert signature_matches("txt", b"just words") is True

    def test_is_case_insensitive_about_the_declared_type(self):
        assert signature_matches("PDF", b"%PDF-1.4") is True
